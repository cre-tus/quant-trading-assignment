"""PPT 템플릿 테스트 (표지·거래안내가 실제 실행을 반영하는지)."""
import unittest
from pathlib import Path
import tempfile

import pandas as pd

from src.report import build_pptx, equity_chart


def _cfg():
    return {
        "universe": ["AAPL"],
        "cash_yield_annual": 0.04,
        "initial_capital": 1000000000,
        "currency": "USD",
        "strategy": {"module": "my_strategy", "rebalance_freq": "2W",
                     "rsi_window": 14, "rsi_sell": 85, "rsi_buy": 20,
                     "top_n": 4, "weight_each": 0.20, "cash_weight": 0.20,
                     "rebalance_band": 0.02,
                     "short": {"signal": "rsi", "rsi_buy": 20, "rsi_sell": 70,
                               "cooldown_days": 10, "stop_loss": 0.10}},
    }


def _metrics():
    return {"CAGR": 0.42, "MDD": -0.17, "MDD_peak": "2024-01-01",
            "MDD_trough": "2024-02-01", "total_return": 1.0,
            "sharpe": 1.39, "volatility": 0.2}


class TestReport(unittest.TestCase):
    def _texts(self, path):
        from pptx import Presentation
        out = []
        for s in Presentation(str(path)).slides:
            out.append(" ".join(sh.text for sh in s.shapes
                                if hasattr(sh, "text")))
        return out

    def _chart(self, td: Path) -> str:
        idx = pd.date_range("2024-01-01", periods=10, freq="B")
        eq = pd.Series(range(10), index=idx, dtype=float) + 100.0
        path = str(td / "equity_curve.png")
        equity_chart(eq, eq, path)
        return path

    def test_cover_and_records_reflect_run(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            out = tdp / "report.pptx"
            fund = pd.DataFrame({"trailingPE": [10.0], "priceToBook": [1.0]},
                                index=["A"])
            build_pptx(_cfg(), _metrics(), 0.27, fund, "edgar", "pool",
                       1800, self._chart(tdp), str(out), ["- test"],
                       universe_info="pool 3175→유동성 3175",
                       outdir="outputs_final_pool_k",
                       strat_name="my_strategy")
            texts = self._texts(out)
            cover = texts[0]
            self.assertIn("3175", cover)
            self.assertNotIn("AAPL", cover)
            self.assertIn("손절 10%", cover)
            last = texts[-1]
            self.assertIn("outputs_final_pool_k/trade_history_view.csv", last)
            self.assertNotIn("outputs/trade_log.csv", last)

    def test_legacy_strategy_fallback(self):
        with tempfile.TemporaryDirectory() as td:
            tdp = Path(td)
            out = tdp / "report.pptx"
            cfg = _cfg()
            cfg["strategy"]["module"] = "example_value_rsi"
            fund = pd.DataFrame({"trailingPE": [10.0], "priceToBook": [1.0]},
                                index=["A"])
            build_pptx(cfg, _metrics(), 0.27, fund, "demo", "demo",
                       10, self._chart(tdp), str(out), ["- test"],
                       strat_name="example_value_rsi")
            texts = self._texts(out)
            self.assertIn("유니버스(10): AAPL", texts[0])
            self.assertIn("전략 옵션1", texts[0])


if __name__ == "__main__":
    unittest.main()
