"""차트 + PPT 보고서 자동생성 (제출용)."""
from __future__ import annotations

from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd


def equity_chart(eq: pd.Series, bh: pd.Series, path: str):
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.plot(eq.index, eq.values, label="Strategy")
    ax.plot(bh.index, bh.values, label="Buy & Hold (equal weight)", linestyle="--")
    ax.set_title("Equity curve (2024-2025)")
    ax.set_xlabel("Date")
    ax.set_ylabel("Portfolio value")
    ax.legend()
    ax.grid(alpha=0.3)
    fig.tight_layout()
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.close(fig)


def _auto_analysis(m: dict, bh_ret: float) -> list[str]:
    good, bad = [], []
    (good if m["CAGR"] > bh_ret / 2 else bad).append(
        f"CAGR {m['CAGR']:.1%} vs B&H 연환산 약 {bh_ret/2:.1%}: "
        + ("초과" if m["CAGR"] > bh_ret / 2 else "미달"))
    (good if m["MDD"] > -0.25 else bad).append(f"MDD {m['MDD']:.1%}")
    (good if m["sharpe"] > 0.8 else bad).append(f"Sharpe {m['sharpe']:.2f}")
    lines = ["[좋았던 점]"]
    lines += [f"- {g}" for g in good] or ["- (두드러진 강점 없음)"]
    lines.append("[나빴던 점 / 원인 가설]")
    lines += [f"- {b}" for b in bad] or ["- (두드러진 약점 없음)"]
    lines.append("- RSI 85/20 잦은 발동 시 회전율·수수료 상승 가능 → fee 민감도 분석 권장")
    lines.append("- 스냅샷 펀더멘털의 사후편향 한계: 실제 제출 전 point-in-time DB로 교체 검토")
    return lines


def build_pptx(cfg: dict, m: dict, bh_ret: float, fund: pd.DataFrame,
               fund_src: str, price_src: str, n_trades: int,
               chart_path: str, out_path: str, analysis_lines: list | None = None,
               universe_info: str = "", outdir: str = "outputs",
               strat_name: str = ""):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    def add(title: str, bullets: list[str]):
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.clear()
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.level = 0

    s = cfg["strategy"]
    sh = s.get("short", {})
    if strat_name.startswith("my_strategy"):
        strat_line = ("스윙 50% + 단기 40% + 현금 10%: PER+PBR 상위 9개 각 10%, "
                      f"{s['rebalance_freq']} 리밸런싱")
    else:
        strat_line = (f"전략 옵션1: PER+PBR 합산랭크 상위 {s['top_n']}개 각 "
                      f"{s['weight_each']:.0%} + 현금 {s['cash_weight']:.0%}")
    sig_line = (f"단기신호 {sh.get('signal', 'rsi')}: 매수 RSI<{sh.get('rsi_buy', s['rsi_buy'])} / "
                f"매도 RSI>{sh.get('rsi_sell', s['rsi_sell'])}, 쿨다운 {sh.get('cooldown_days', 0)}일")
    if float(sh.get("stop_loss", 0.0) or 0.0) > 0:
        sig_line += f", 손절 {float(sh['stop_loss']):.0%}"
    sig_line += f" / 리밸런싱 밴드 {float(s.get('rebalance_band', 0.0)):.2%}"
    uni_line = (f"유니버스: {universe_info}" if universe_info
                else f"유니버스(10): {', '.join(cfg['universe'])}")
    add("리밸런싱 백테스트 결과 (2024.1.1~2025.12.31)", [
        f"{uni_line} + 현금(MMF {cfg['cash_yield_annual']:.0%})",
        strat_line,
        sig_line,
        f"데이터: 가격={price_src}, 펀더멘털={fund_src} / 초기 {cfg['initial_capital']:,} {cfg['currency']}",
    ])
    add("개발 도구 및 방식", [
        "Python 백테스트 엔진(src/ + run_backtest.py, config.yaml 기반 재현)",
        "데이터: yfinance(일봉 시세) + SEC EDGAR(펀더멘털 확정치) + NASDAQ 상장명단",
        "AI 활용: Muse Code(Muse Spark)로 엔진 생성, 프롬프트 prompts/ 보관(제출물 2-1)",
        "산출물: trade_log / daily_holdings / discovery / trade_history_view.csv, "
        "daily_equity.csv, metrics.json, report.xlsx",
    ])
    add("핵심 성과지표 (필수 3종 포함)", [
        f"PL CAGR(전략): {m['CAGR']:.2%}",
        f"Buy & Hold Return(균등비중): {bh_ret:.2%}",
        f"Max Drawdown: {m['MDD']:.2%} ({m['MDD_peak']} → {m['MDD_trough']})",
        f"총수익률 {m['total_return']:.2%} / Sharpe {m['sharpe']:.2f} / 변동성 {m['volatility']:.2%} / 총 거래 {n_trades}건",
    ])
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = "누적 수익률 곡선"
    sl.shapes.add_picture(chart_path, Inches(0.5), Inches(1.2), width=Inches(12.3), height=Inches(5.8))
    show = fund.head(25)
    rows = []
    for idx, row in show.iterrows():
        pe = f"{row['trailingPE']:.1f}" if pd.notna(row["trailingPE"]) else "결측"
        pb = f"{row['priceToBook']:.1f}" if pd.notna(row["priceToBook"]) else "결측"
        rows.append(f"{idx}: PE {pe} / PB {pb}")
    if len(fund) > 25:
        rows.append(f"... 외 {len(fund) - 25}종목 (data/fundamentals_snapshot.csv 참조)")
    add(f"펀더멘털 스냅샷 (PER·PBR 랭킹 근거, {len(fund)}종목 중 25)", rows)
    add("결과 분석 및 평가", analysis_lines or _auto_analysis(m, bh_ret))
    add("거래 기록 안내", [
        f"매 거래 시점: {outdir}/trade_history_view.csv (순번·매수/매도·종목·수량·가격·실현손익·거래후 자산)",
        f"원장: {outdir}/trade_log.csv, 일별보유: {outdir}/daily_holdings.csv, 종목손익: {outdir}/discovery.csv",
        f"지표: {outdir}/metrics.json·analysis.md, 엑셀: {outdir}/report.xlsx (원본 CSV 함께 제출)",
    ])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
