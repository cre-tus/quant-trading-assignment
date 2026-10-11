"""산출물 엑셀 변환 (#### 방지: 날짜 형식 + 열 너비 + 틀 고정 + 필터).

CSV를 엑셀로 열 때 날짜가 ####로 보이는 문제를 없애기 위해, 파이프라인과
make_excel.py가 매 실행 폴더에 report.xlsx를 만든다.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd

SHEETS = (("daily_holdings.csv", "일별보유"),
          ("discovery.csv", "종목손익"),
          ("trade_log.csv", "거래기록"))

DATE_FMT = "yyyy-mm-dd"


def _fit_sheet(ws, df: pd.DataFrame, date_cols: set[str]):
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(df.columns, start=1):
        cell_col = ws.column_dimensions[ws.cell(row=1, column=j).column_letter]
        if col in date_cols:
            cell_col.width = 13
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                row[0].number_format = DATE_FMT
            continue
        try:
            longest = max([len(str(col))] +
                          [len(str(v)) for v in df[col].astype(str).head(500)])
        except Exception:
            longest = len(str(col))
        cell_col.width = min(max(longest + 2, 12), 28)


def build_workbook(outdir: str | Path) -> str:
    """outdir 안 CSV들로 report.xlsx 생성. 경로 반환."""
    from openpyxl import Workbook
    odir = Path(outdir)
    wb = Workbook()
    wb.remove(wb.active)
    made = []
    for fname, title in SHEETS:
        p = odir / fname
        if not p.exists():
            continue
        df = pd.read_csv(p)
        ws = wb.create_sheet(title)
        ws.append(list(df.columns))
        date_cols = {c for c in df.columns if "date" in c.lower()}
        for _, row in df.iterrows():
            vals = []
            for c in df.columns:
                v = row[c]
                if c in date_cols:
                    try:
                        v = pd.Timestamp(v).to_pydatetime()
                    except (ValueError, TypeError):
                        pass
                vals.append(v)
            ws.append(vals)
        _fit_sheet(ws, df, date_cols)
        made.append(title)
    mpath = odir / "metrics.json"
    if mpath.exists():
        import json
        m = json.loads(mpath.read_text(encoding="utf-8"))
        ws = wb.create_sheet("지표")
        s = m.get("strategy", {})
        for k in ("CAGR", "total_return", "MDD", "sharpe", "volatility"):
            ws.append([k, s.get(k)])
        ws.append(["buy_and_hold_return", m.get("buy_and_hold_return")])
        ws.append(["n_trades", m.get("n_trades")])
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 18
        made.append("지표")
    if not made:
        raise ValueError(f"변환할 CSV가 없음: {odir}")
    out = str(odir / "report.xlsx")
    wb.save(out)
    return out
